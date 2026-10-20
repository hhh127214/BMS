# -*- coding: utf-8 -*-
"""以参考说明书为版面底本，生成 BMS 平台操作说明书。

做法：直接打开参考 docx 当模板（继承 styles.xml / numbering.xml / 页面设置），
清空正文后按块列表重建内容 —— 这样字体、标题层级、多级编号、表格样式
与参考文档完全一致，而不是"另起一套模板"。
"""
import copy
import os
import re
import sys

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING,
                            WD_TAB_ALIGNMENT, WD_TAB_LEADER)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEMPLATE = (r'C:\Users\31857\Documents\xwechat_files\wxid_z3jddcym371q22_58a6'
            r'\msg\file\2026-10'
            r'\03 20260818_EMS-SCADA控制软件V1.0操作说明书 递交版_唐政反馈(2)(1).docx')
SHOT = r'D:\wb(cn)\BMS\docs\manual\shots'

USABLE_CM = 17.19          # A4 21cm - 左右各 1.905cm
FIG_W_CM = 16.6


# ---------------------------------------------------------------- 基础工具
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag == qn('w:sectPr'):
            continue
        body.remove(child)


def setup_page(doc):
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(1.905)
        sec.right_margin = Cm(1.905)
        sec.top_margin = Cm(1.905)
        sec.bottom_margin = Cm(1.905)
        sec.header_distance = Cm(1.2)
        sec.footer_distance = Cm(1.1)
        # 模板沿用 numberInDash（- 1 -）；页脚已手工加破折号，故改回纯数字
        sectPr = sec._sectPr
        pgNum = sectPr.find(qn('w:pgNumType'))
        if pgNum is None:
            pgNum = _el('w:pgNumType')
            sectPr.append(pgNum)
        pgNum.set(qn('w:fmt'), 'decimal')


def keep_one_section(doc):
    """把模板里的 5 节压成 1 节：只保留最后一个 sectPr，其余连同分节段落清掉。"""
    body = doc.element.body
    sect = [c for c in body if c.tag == qn('w:sectPr')]
    for s in sect[1:]:
        body.remove(s)
    # 段落级 sectPr 在 clear_body 时已随段落删除
    return doc.sections[0]


def enable_update_fields(doc):
    """打开文档时自动更新域（目录、页码）。"""
    st = doc.settings.element
    if st.find(qn('w:updateFields')) is None:
        st.append(_el('w:updateFields', **{'w:val': 'true'}))


def add_field(paragraph, instr, placeholder='1'):
    fld = _el('w:fldSimple', **{'w:instr': instr})
    r = _el('w:r')
    t = _el('w:t')
    t.text = placeholder
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


# ---------------------------------------------------------------- 段落
_BOLD = re.compile(r'\*\*(.+?)\*\*')


def _add_runs(par, text, size=None, bold_all=False, color=None, font=None):
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            _style_run(r, size, bold_all, color, font)
        r = par.add_run(m.group(1))
        _style_run(r, size, True, color, font)
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:])
        _style_run(r, size, bold_all, color, font)


def _style_run(run, size, bold, color, font):
    if size:
        run.font.size = Pt(size)
    if bold:
        run.font.bold = True
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if font:
        run.font.name = font
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font)


def P(doc, text, style='Normal', align=None, indent=True, size=None,
      bold=False, space_before=None, space_after=None, color=None):
    par = doc.add_paragraph(style=style)
    if not indent:
        par.paragraph_format.first_line_indent = Cm(0)
    if align == 'center':
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'right':
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == 'justify':
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if space_before is not None:
        par.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    if text:
        _add_runs(par, text, size=size, bold_all=bold, color=color)
    return par


def H(doc, level, text):
    return doc.add_paragraph(text, style='Heading %d' % level)


def H1N(doc, text):
    """一级标题但不参与自动编号（附录用）。"""
    par = doc.add_paragraph(text, style='Heading 1')
    pPr = par._p.get_or_add_pPr()
    try:
        numPr = pPr.get_or_add_numPr()
        numPr.get_or_add_ilvl().val = 0
        numPr.get_or_add_numId().val = 0
    except Exception:
        numPr = _el('w:numPr')
        numPr.append(_el('w:ilvl', **{'w:val': '0'}))
        numPr.append(_el('w:numId', **{'w:val': '0'}))
        pPr.append(numPr)
    return par


def NOTE(doc, text):
    """参考文档的「注意：xxx」一行 —— 加粗引导 + 常缩进正文。"""
    par = doc.add_paragraph(style='Normal')
    r = par.add_run('注意：')
    r.font.bold = True
    _add_runs(par, text)
    return par


def BULLET(doc, text, level=0):
    par = doc.add_paragraph(style='Normal')
    pf = par.paragraph_format
    pf.left_indent = Cm(0.74 + 0.74 * level)
    pf.first_line_indent = Cm(0)
    r = par.add_run('· ')
    _add_runs(par, text)
    return par


def PAGEBREAK(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- 图
def FIG(doc, filename, caption=None):
    path = os.path.join(SHOT, filename)
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.first_line_indent = Cm(0)
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.line_spacing = 1.0
    par.add_run().add_picture(path, width=Cm(FIG_W_CM))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.first_line_indent = Cm(0)
        cap.paragraph_format.space_after = Pt(8)
        cap.paragraph_format.line_spacing = 1.0
        r = cap.add_run(caption)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string('595959')


# ---------------------------------------------------------------- 表格
def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(_el('w:shd', **{'w:val': 'clear', 'w:color': 'auto',
                                'w:fill': hexcolor}))


def _cell_text(cell, text, size=10.5, bold=False, align=None, color=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell.text = ''
    par = cell.paragraphs[0]
    par.paragraph_format.first_line_indent = Cm(0)
    par.paragraph_format.space_before = Pt(1)
    par.paragraph_format.space_after = Pt(1)
    par.paragraph_format.line_spacing = 1.0
    if align == 'center':
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'justify':
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_runs(par, text, size=size, bold_all=bold, color=color)



def _tbl_finish(t, left=76, right=76, top=20, bottom=20):
    """收窄单元格左右内边距（默认 0.19cm 会把四字词挤成两行），并去掉重复的 tblLayout。"""
    tblPr = t._tbl.tblPr
    for dup in tblPr.findall(qn('w:tblLayout'))[1:]:
        tblPr.remove(dup)
    mar = OxmlElement('w:tblCellMar')
    for tag, val in (('w:top', top), ('w:left', left),
                     ('w:bottom', bottom), ('w:right', right)):
        mar.append(_el(tag, **{'w:w': val, 'w:type': 'dxa'}))
    look = tblPr.find(qn('w:tblLook'))
    if look is not None:
        tblPr.remove(look)
    old = tblPr.find(qn('w:tblCellMar'))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(mar)
    if look is not None:
        tblPr.append(look)


def TABLE(doc, headers, rows, widths=None, size=10.5, align_center=None,
          header_fill='EEF3F9'):
    """align_center: 需要居中的列序号集合。"""
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    _tbl_finish(t)
    if widths:
        total = sum(widths)
        widths = [w * USABLE_CM / total for w in widths]
        # 保底列宽：过窄的列会把「10」「11」这类两位数竖排成两行
        MIN = 2.0
        for i, w in enumerate(widths):
            if w < MIN:
                j = widths.index(max(widths))
                widths[j] = widths[j] - (MIN - w)
                widths[i] = MIN
    else:
        widths = [USABLE_CM / len(headers)] * len(headers)
    for idx, w in enumerate(widths):
        t.columns[idx].width = Cm(w)
    for row in t.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Cm(widths[idx])
    align_center = align_center or set()
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(_el('w:tblHeader', **{'w:val': 'true'}))
    for idx, h in enumerate(headers):
        c = t.rows[0].cells[idx]
        _cell_text(c, h, size=size, bold=True,
                   align='center' if True else None)
        _shade(c, header_fill)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.rows[ri + 1].cells[ci]
            _cell_text(c, str(val), size=size,
                       align='center' if ci in align_center else None)
    # 表后留白
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(4)
    tail.paragraph_format.line_spacing = 1.0
    tail.paragraph_format.first_line_indent = Cm(0)
    # ★ 防孤行：表头与首条数据行不得单独落在页底（keepNext 连成一组，整组下移）。
    #   只锁这两行，长表仍可在后续行处自然分页。
    for r in ([t.rows[0]] + ([t.rows[1]] if len(rows) >= 2 else [])):
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    # 单元格内容不得跨页拆开
    for r in t.rows:
        r._tr.get_or_add_trPr().append(_el('w:cantSplit', **{'w:val': 'true'}))
    return t


def CODE(doc, text, size=9):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _tbl_finish(t)
    t.columns[0].width = Cm(USABLE_CM)
    cell = t.rows[0].cells[0]
    cell.width = Cm(USABLE_CM)
    _shade(cell, 'F4F6F8')
    cell.text = ''
    lines = text.split('\n')
    for i, line in enumerate(lines):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        pf = par.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        r = par.add_run(line if line else ' ')
        r.font.size = Pt(size)
        r.font.name = 'Consolas'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(4)
    tail.paragraph_format.line_spacing = 1.0
    tail.paragraph_format.first_line_indent = Cm(0)


# ---------------------------------------------------------------- 结构件
def cover(doc, title, subtitle, lines):
    for _ in range(4):
        par = doc.add_paragraph()
        par.paragraph_format.line_spacing = 1.0
    for line in title.split('\n'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(line)
        r.font.size = Pt(30)
        r.font.bold = True
        r.font.name = '黑体'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run(subtitle)
    r.font.size = Pt(18)
    r.font.name = '黑体'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    for _ in range(6):
        par = doc.add_paragraph()
        par.paragraph_format.line_spacing = 1.0

    for k, v in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run('%s：%s' % (k, v))
        r.font.size = Pt(12)
        r.font.name = '宋体'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def _mulu_styles(doc):
    """目录专用样式（自定义名，避开 Word 自带的 TOC 样式自动重定义）。"""
    for name, indent, before, bold, ea in (
            ('Mulu 1', 0.00, 7, True, '黑体'),
            ('Mulu 2', 0.85, 0, False, '仿宋')):
        try:
            st = doc.styles[name]
        except KeyError:
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = doc.styles['Normal']
        pf = st.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(indent)
        pf.space_before = Pt(before)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.2
        pf.tab_stops.add_tab_stop(Cm(USABLE_CM), WD_TAB_ALIGNMENT.RIGHT,
                                  WD_TAB_LEADER.DOTS)
        f = st.font
        f.name = 'Times New Roman'
        f.size = Pt(11)
        f.bold = bold
        rPr = st.element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = _el('w:rFonts')
            rPr.append(rFonts)
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rFonts.set(qn('w:eastAsia'), ea)


def toc_entries(blocks):
    """按块列表推算目录条目：h1 记章号，h2 记节号，h1n（附录）无编号且不占章号。"""
    out = []
    ch = 0
    sec = 0
    for b in blocks:
        if b[0] == 'h1':
            ch += 1
            sec = 0
            out.append((1, '%d.' % ch, b[1]))
        elif b[0] == 'h2':
            sec += 1
            out.append((2, '%d.%d.' % (ch, sec), b[1]))
        elif b[0] == 'h1n':
            out.append((1, '', b[1]))
    return out


def toc(doc, entries, pages=None):
    """静态目录：条目与页码直接落成文字，右对齐点线制表位，页码由 Word 分页结果回填。"""
    _mulu_styles(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run('目　　录')
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.name = '黑体'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    for idx, (lvl, num, title) in enumerate(entries):
        par = doc.add_paragraph(style='Mulu 1' if lvl == 1 else 'Mulu 2')
        if num:
            par.add_run(num + '\u3000')
        par.add_run(title)
        tab = par.add_run()
        tab.add_tab()
        # 页码：构建前两轮——先占位（与最终等宽），取 Word 分页结果后再写实数
        par.add_run(str(pages[idx]) if pages else '00')


def footer_pageno(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    sec.footer.is_linked_to_previous = False
    par = sec.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.first_line_indent = Cm(0)
    r = par.add_run('— ')
    r.font.size = Pt(9)
    add_field(par, ' PAGE ', '1')
    r2 = par.add_run(' —')
    r2.font.size = Pt(9)


def build(blocks, out_path, pages=None):
    doc = docx.Document(TEMPLATE)
    clear_body(doc)
    setup_page(doc)
    keep_one_section(doc)
    footer_pageno(doc)
    entries = toc_entries(blocks)

    for b in blocks:
        kind = b[0]
        if kind == 'h1':
            H(doc, 1, b[1])
        elif kind == 'h1n':
            H1N(doc, b[1])
        elif kind == 'cover':
            cover(doc, *b[1:])
        elif kind == 'toc':
            toc(doc, entries, pages)
        elif kind == 'h2':
            H(doc, 2, b[1])
        elif kind == 'h3':
            H(doc, 3, b[1])
        elif kind == 'p':
            P(doc, b[1])
        elif kind == 'pc':                     # 无缩进正文
            P(doc, b[1], indent=False)
        elif kind == 'note':
            NOTE(doc, b[1])
        elif kind == 'b':
            BULLET(doc, b[1], b[2] if len(b) > 2 else 0)
        elif kind == 'img':
            FIG(doc, b[1], b[2] if len(b) > 2 else None)
        elif kind == 'table':
            widths = b[3] if len(b) > 3 else None
            size = b[4] if len(b) > 4 else 10.5
            if isinstance(size, (list, tuple)):
                size = size[0]
            ac = set(b[5]) if len(b) > 5 else None
            TABLE(doc, b[1], b[2], widths, size, ac)
        elif kind == 'code':
            CODE(doc, b[1])
        elif kind == 'pb':
            PAGEBREAK(doc)
        else:
            raise ValueError('unknown block %r' % (kind,))

    doc.save(out_path)
    return out_path


if __name__ == '__main__':
    import importlib
    mod = importlib.import_module(sys.argv[1])
    print(build(mod.BLOCKS, sys.argv[2]))
